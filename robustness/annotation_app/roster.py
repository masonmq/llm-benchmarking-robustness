"""Manage annotator usernames and paper assignments.

The roster spreadsheet identifies people, so it lives under robustness/data/ (gitignored).
Usernames starting with `cr-` review code (Part 3); `an-` usernames skip Part 3.

  python roster.py generate --reviewers 15 --annotators 15    # create or extend the roster
  python roster.py allowed                                    # value for the ALLOWED_ANNOTATORS secret
  python roster.py assignments                                # SQL that syncs D1 assignments from the roster
"""

from __future__ import annotations

import argparse
import secrets
from datetime import date
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.worksheet.datavalidation import DataValidation

APP_DIR = Path(__file__).resolve().parent
DEFAULT_ROSTER = APP_DIR.parent / "data" / "annotators" / "annotator_roster.xlsx"
COLUMNS = ["username", "role", "assigned_papers", "name", "email", "date_issued", "status", "notes"]
ROLES = {"cr-": "code_reviewer", "an-": "annotator"}
# Omits look-alike characters (0/o, 1/l/i) so usernames survive being read aloud or retyped.
ALPHABET = "abcdefghjkmnpqrstuvwxyz23456789"


def read_rows(path: Path) -> list[dict]:
    sheet = load_workbook(path).active
    header = [cell.value for cell in sheet[1]]
    return [dict(zip(header, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]


def write_rows(path: Path, rows: list[dict]) -> None:
    book = Workbook()
    sheet = book.active
    sheet.title = "roster"
    sheet.append(COLUMNS)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    for row in rows:
        sheet.append([row.get(column) for column in COLUMNS])
    status = DataValidation(type="list", formula1='"available,issued,active,completed,withdrawn"', allow_blank=True)
    sheet.add_data_validation(status)
    status.add(f"G2:G{max(len(rows) + 1, 2)}")
    for column, width in zip("ABCDEFGH", (14, 15, 45, 24, 30, 12, 12, 30)):
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    path.parent.mkdir(parents=True, exist_ok=True)
    book.save(path)


def new_username(prefix: str, taken: set[str]) -> str:
    while True:
        name = prefix + "".join(secrets.choice(ALPHABET) for _ in range(6))
        if name not in taken:
            taken.add(name)
            return name


def papers(row: dict) -> list[str]:
    return [paper.strip() for paper in str(row.get("assigned_papers") or "").replace(";", ",").split(",") if paper.strip()]


def generate(args) -> None:
    rows = read_rows(args.roster) if args.roster.exists() else []
    taken = {str(row["username"]).casefold() for row in rows} | {name.casefold() for name in args.reserved}
    for prefix, count in (("cr-", args.reviewers), ("an-", args.annotators)):
        rows += [{"username": new_username(prefix, taken), "role": ROLES[prefix], "status": "available"} for _ in range(count)]
    write_rows(args.roster, rows)
    print(f"{len(rows)} usernames in {args.roster}")


def allowed(args) -> None:
    names = [str(row["username"]) for row in read_rows(args.roster) if row.get("status") != "withdrawn"]
    print(",".join(dict.fromkeys(names + args.extra)))


def assignments(args) -> None:
    rows = read_rows(args.roster)
    # Lowercase so the worker's case-insensitive check can use the primary-key index.
    quote = lambda value: "'" + str(value).lower().replace("'", "''") + "'"
    lines = ["DELETE FROM assignments;"]
    lines += [f"INSERT INTO assignments(annotator, paper_id) VALUES({quote(row['username'])}, {quote(paper)});"
              for row in rows if row.get("status") != "withdrawn" for paper in papers(row)]
    lines += [f"INSERT INTO assignments(annotator, paper_id) VALUES({quote(name)}, {quote(paper)});"
              for name, paper in (item.split(":", 1) for item in args.extra)]
    # Staff accounts (e.g. the dev team) get every paper that appears in the roster.
    all_papers = sorted({paper for row in rows if row.get("status") != "withdrawn" for paper in papers(row)})
    lines += [f"INSERT INTO assignments(annotator, paper_id) VALUES({quote(name)}, {quote(paper)});"
              for name in args.staff for paper in all_papers]
    args.output.write_text("\n".join(lines) + "\n", encoding="utf-8")
    # --command, not --file: wrangler's file import endpoint rejects OAuth logins.
    print(f"Wrote {len(lines) - 1} assignments to {args.output}. Apply from deployment/cloudflare with:\n"
          f'  npx wrangler d1 execute robustness-annotations --remote --command "$(cat {args.output})"')


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--roster", type=Path, default=DEFAULT_ROSTER)
    commands = parser.add_subparsers(required=True)

    command = commands.add_parser("generate", help="Add new random usernames to the roster")
    command.add_argument("--reviewers", type=int, default=0, help="Number of cr- (code reviewer) usernames")
    command.add_argument("--annotators", type=int, default=0, help="Number of an- usernames")
    command.add_argument("--reserved", nargs="*", default=[], help="Existing usernames that must not be reused")
    command.set_defaults(run=generate)

    command = commands.add_parser("allowed", help="Print the ALLOWED_ANNOTATORS value")
    command.add_argument("--extra", nargs="*", default=[], help="Usernames outside the roster to keep, e.g. pilot accounts")
    command.set_defaults(run=allowed)

    command = commands.add_parser("assignments", help="Write SQL that replaces the D1 assignments table")
    command.add_argument("--output", type=Path, default=DEFAULT_ROSTER.with_name("assignments.sql"))
    command.add_argument("--extra", nargs="*", default=[], help="USERNAME:PAPER_ID pairs outside the roster")
    command.add_argument("--staff", nargs="*", default=[], help="Usernames given access to every paper in the roster")
    command.set_defaults(run=assignments)

    args = parser.parse_args()
    args.run(args)


if __name__ == "__main__":
    main()
