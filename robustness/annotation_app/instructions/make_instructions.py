"""Fill the instruction templates for one week's assignments and render each as a PDF.

  python make_instructions.py --week 1 --deadline "Friday, October 9, 2026"

Reads the "Round 1 - by paper" tab of the assignment workbook, writes one Markdown and one PDF per
person to robustness/data/annotators/week<N>_instructions/ (gitignored). PDFs are printed with Chrome.
"""

from __future__ import annotations

import argparse
import json
import re
import subprocess
from pathlib import Path

import markdown
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ANNOTATORS = HERE.parents[1] / "data" / "annotators"
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"
# Site link and contact emails are private: kept in local_accounts.json (not in git).
_LOCAL = json.loads((ANNOTATORS / "local_accounts.json").read_text(encoding="utf-8"))
SITE = _LOCAL["site"]
CONTACT = _LOCAL["contact"]
CSS = """
@page { size: Letter; margin: 0.8in 0.85in; }
body { font-family: -apple-system, "Helvetica Neue", Arial, sans-serif; font-size: 10.5pt; line-height: 1.5; color: #1f2933; }
h1 { font-size: 19pt; color: #17324d; margin: 0 0 6pt; padding-bottom: 8pt; border-bottom: 2px solid #17324d; }
h2 { font-size: 13.5pt; color: #17324d; margin: 18pt 0 6pt; padding-bottom: 3pt; border-bottom: 1px solid #d5dde5; break-after: avoid; }
h3 { font-size: 11.5pt; color: #24496b; margin: 13pt 0 4pt; break-after: avoid; }
p { margin: 5pt 0; }
ul, ol { margin: 4pt 0 6pt; padding-left: 20pt; }
li { margin: 2.5pt 0; }
a { color: #176b91; }
code { font-family: Menlo, monospace; font-size: 9.5pt; background: #f1f4f7; padding: 0 3pt; border-radius: 3px; }
h2#access + ul { list-style: none; padding: 9pt 12pt; margin: 6pt 0 10pt; background: #f2f6fa; border: 1px solid #cfdbe6; border-radius: 6px; }
li, p { break-inside: avoid; }
"""


def week_assignments(workbook: Path, week: int) -> list[tuple[int, str, str, str]]:
    """(screening ID, username, role, paper_id) for everyone assigned a paper in the given week."""
    sheet = load_workbook(workbook)["Round 1 - by paper"]
    header = [cell.value for cell in sheet[1]]
    people = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        record = dict(zip(header, row))
        if record["Week"] != week:
            continue
        for column in ("Code reviewer 1", "Code reviewer 2", "Annotator(s)"):
            for match in re.finditer(r"ID (\d+) \(([a-z]{2}-[a-z0-9]+)\)", str(record[column] or "")):
                role = "code_reviewer" if match.group(2).startswith("cr-") else "annotator"
                people.append((int(match.group(1)), match.group(2), role, record["Paper ID"]))
    return sorted(people)


def render_pdf(markdown_text: str, pdf_path: Path) -> None:
    body = markdown.markdown(markdown_text, extensions=["toc", "sane_lists"])
    html_path = pdf_path.with_suffix(".html")
    html_path.write_text(f"<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{body}</body></html>", encoding="utf-8")
    try:
        subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_path}", html_path.resolve().as_uri()], check=True, capture_output=True)
    finally:
        html_path.unlink(missing_ok=True)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--week", type=int, required=True)
    parser.add_argument("--deadline", required=True, help='e.g. "Friday, October 9, 2026"')
    parser.add_argument("--workbook", type=Path, default=ANNOTATORS / "round1_assignment.xlsx")
    args = parser.parse_args()

    output = ANNOTATORS / f"week{args.week}_instructions"
    (output / "pdf").mkdir(parents=True, exist_ok=True)
    for screening_id, username, role, paper in week_assignments(args.workbook, args.week):
        text = (HERE / f"INSTRUCTIONS_{role}.md").read_text(encoding="utf-8")
        for key, value in {"WEEK": str(args.week), "PAPER_LINK": SITE.format(paper), "USERNAME": username,
                           "DEADLINE": args.deadline, "CONTACT": CONTACT}.items():
            text = text.replace("{{" + key + "}}", value)
        assert "{{" not in text, f"unfilled placeholder for ID {screening_id}"
        stem = f"Week{args.week}_ID{screening_id:02d}_{username}_{paper.split('_')[0]}"
        (output / f"{stem}.md").write_text(text, encoding="utf-8")
        render_pdf(text, output / "pdf" / f"{stem}.pdf")
        print(f"{stem}.pdf")


if __name__ == "__main__":
    main()
