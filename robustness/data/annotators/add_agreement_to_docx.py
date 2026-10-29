"""Insert the reviewer-agreement table into an existing (hand-edited) report .docx, leaving everything else as is.

  robustness/annotation_app/.venv/bin/python robustness/data/annotators/analyze_agreement.py --weeks 1
  robustness/annotation_app/.venv/bin/python robustness/data/annotators/add_agreement_to_docx.py report/round1_week1_report.docx --weeks 1

Numbers come from report/agreement_week<weeks>.json. Part 2 questions use agreement after recoding to
Yes vs. not Yes; the 1-5 ratings (Parts 3 and 4) use agreement within 1 point. The section goes just before the "Reading these results" heading (or at the end if that
heading is gone). A copy of the original file is saved next to it as <name>.before-agreement.docx.
"""

from __future__ import annotations

import argparse
import json
import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

OUT = Path(__file__).resolve().parent / "report"
HEADING = "Reviewer agreement"


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("docx", type=Path)
    parser.add_argument("--weeks", required=True)
    args = parser.parse_args()
    if (args.docx.parent / f"~${args.docx.name[2:]}").exists():
        raise SystemExit("The document looks open in Word (lock file present). Save and close it first.")

    rows = json.loads((OUT / f"agreement_week{args.weeks.replace(',', '-')}.json").read_text())["overall"]
    doc = Document(args.docx)
    if any(p.text.strip() == HEADING for p in doc.paragraphs):
        raise SystemExit(f"The document already has a \"{HEADING}\" section.")
    shutil.copy2(args.docx, args.docx.with_suffix(".before-agreement.docx"))

    # Build the new section at the end, then move it before "Reading these results" if that heading exists.
    heading = doc.add_heading(HEADING, level=1)
    intro = doc.add_paragraph("Share of reviewer pairs on the same analysis who agreed. Part 2 answers are recoded to Yes vs. not Yes "
                      "(No, Unclear or Insufficient information) before comparing. "
                      "The 1-5 ratings (Parts 3 and 4) count answers within 1 point of each other.")
    t = doc.add_table(rows=1, cols=3)
    t.style = "Table Grid"
    for cell, text in zip(t.rows[0].cells, ["Question", "Analyses", "Agreement"]):
        cell.text = ""
        cell.paragraphs[0].add_run(text).bold = True
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "EDEDED")
        cell._tc.get_or_add_tcPr().append(shd)
    for r in rows:
        question = r["question"]
        if r["level"] == "ordinal":
            share = r["within_1"]
        else:  # Part 2 questions recoded to Yes vs. not Yes (No, Unclear, Insufficient information)
            share = r["exact_agreement_yes_vs_not"]
        cells = t.add_row().cells
        cells[0].text, cells[1].text, cells[2].text = question, str(r["analyses"]), f"{share:.0%}"
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for run in p.runs:
                    run.font.size = Pt(10)
    spacer = doc.add_paragraph()
    new = [heading._p, intro._p, t._tbl, spacer._p]

    anchor = next((p for p in doc.paragraphs if p.text.strip() == "Reading these results"), None)
    if anchor is not None:
        for el in new:
            anchor._p.addprevious(el)
    doc.save(args.docx)
    print(f"Added \"{HEADING}\" to {args.docx} ({'before \"Reading these results\"' if anchor is not None else 'at the end'}).")


if __name__ == "__main__":
    main()
